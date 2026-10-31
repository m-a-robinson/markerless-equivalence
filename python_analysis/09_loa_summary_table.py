#!/usr/bin/env python3
"""Manuscript table: functional limits of agreement (LoA) between markerless and marker-based kinematics.

Reads ``results/loa_<tag>.npz`` written by ``03_loa_analysis.py``.  For every variable, the pointwise functional LoA
(mean difference MB - ML +/- 1.96 SD of the participant differences at each point of the gait cycle; Roislien et al. 2012)
are summarised over the 101 gait-cycle points:

* mean difference    - average over the cycle, with the range (minimum to maximum) in brackets;
* LoA width          - upper minus lower LoA, averaged over the cycle, with the maximum in brackets;
* % of the cycle in each LoA outcome, judged against the bounds (``common.classify_loa``):
    acceptable      the LoA lie inside the bounds (individual differences expected to be acceptable);
    group level only  the LoA exceed a bound but the mean difference is inside;
    systematic bias the mean difference itself is beyond a bound.

Rows are ordered by % acceptable (then by narrower mean width) and the last row is the mean over the variables.

Outputs (``results/``): ``table_loa_<tag>.csv`` (all numbers) and, if python-docx is installed, ``table_loa_<tag>.docx``
(with merged headers and a footnote).

Usage
-----
    python 09_loa_summary_table.py [--tag 5deg]
"""
import argparse
import csv
import json
from pathlib import Path

import numpy as np

from common import LOA_ACCEPTABLE, LOA_BIAS, LOA_GROUP_ONLY, RESULTS_DIR, classify_loa

OUTCOMES = [('acceptable', LOA_ACCEPTABLE), ('group_only', LOA_GROUP_ONLY), ('bias', LOA_BIAS)]


def build_rows(z, meta):
    bound = meta['bound']
    rows = []
    for code, title, n in zip(meta['variables'], meta['titles'], meta['n']):
        md, lo, hi = z[f'{code}_md'], z[f'{code}_lo'], z[f'{code}_hi']
        w, cls = hi - lo, classify_loa(md, lo, hi, bound)
        r = dict(variable=title, n=int(n), bias_mean=float(md.mean()), bias_min=float(md.min()), bias_max=float(md.max()),
                 width_mean=float(w.mean()), width_max=float(w.max()))
        for name, c in OUTCOMES:
            r[f'pct_{name}'] = 100 * float(np.mean(cls == c))
        rows.append(r)
    rows.sort(key=lambda r: (-r['pct_acceptable'], r['width_mean']))
    keys = [k for k in rows[0] if k not in ('variable', 'n')]
    rows.append(dict(variable='Mean of the variables', n='', **{k: float(np.mean([r[k] for r in rows])) for k in keys}))
    return rows, bound


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--tag', default='5deg', help='tag of the LoA results (default 5deg)')
    ap.add_argument('--results-dir', type=Path, default=RESULTS_DIR)
    args = ap.parse_args()

    z = np.load(args.results_dir / f'loa_{args.tag}.npz', allow_pickle=False)
    meta = json.loads(str(z['meta']))
    rows, bound = build_rows(z, meta)

    with open(args.results_dir / f'table_loa_{args.tag}.csv', 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(['Variable', 'n', 'MeanDifference_mean_deg', 'MeanDifference_min_deg', 'MeanDifference_max_deg',
                    'LoAwidth_mean_deg', 'LoAwidth_max_deg', 'Acceptable_pct', 'GroupLevelOnly_pct', 'SystematicBias_pct'])
        for r in rows:
            mean_row = r['n'] == ''                                  # the last row: only the cycle averages are meaningful
            w.writerow([r['variable'], r['n'], '' if mean_row else f"{r['bias_mean']:.2f}", '' if mean_row else f"{r['bias_min']:.2f}",
                        '' if mean_row else f"{r['bias_max']:.2f}", f"{r['width_mean']:.2f}", '' if mean_row else f"{r['width_max']:.2f}",
                        f"{r['pct_acceptable']:.1f}",
                        f"{r['pct_group_only']:.1f}", f"{r['pct_bias']:.1f}"])

    def diff_txt(r):
        return '–' if r['n'] == '' else f"{r['bias_mean']:.1f} [{r['bias_min']:.1f}, {r['bias_max']:.1f}]"   # no mean of signed values

    def width_txt(r):
        return f"{r['width_mean']:.1f}" if r['n'] == '' else f"{r['width_mean']:.1f} [{r['width_max']:.1f}]"

    footnote = (f'Pointwise functional limits of agreement (LoA): mean difference (marker-based − markerless) ± 1.96 SD of the '
                f'participant differences at each of 101 points of the gait cycle. Mean difference: average over the cycle '
                f'[minimum, maximum]. LoA width (upper − lower LoA): average over the cycle [maximum]. Outcomes are the '
                f'percentage of the gait cycle with the LoA inside ±{bound:g}° (acceptable), the LoA exceeding a bound but the '
                f'mean difference inside ±{bound:g}° (group level only), or the mean difference beyond ±{bound:g}° (systematic '
                f'bias). Percentages are rounded.')
    print(f'LoA summary (bound ±{bound:g}°)')
    print(f"{'Variable':22s}{'n':>3s}  {'Mean diff [min, max]':>24s}  {'Width [max]':>14s}  accept./group/bias %")
    for r in rows:
        print(f"{r['variable']:22s}{str(r['n']):>3s}  {diff_txt(r):>24s}  {width_txt(r):>14s}  "
              f"{r['pct_acceptable']:5.0f} / {r['pct_group_only']:3.0f} / {r['pct_bias']:3.0f}")

    try:
        import docx
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
    except ImportError:
        print('python-docx not installed - table written as .csv only')
        return
    d = docx.Document()
    ncol = 7
    t = d.add_table(rows=2, cols=ncol)
    t.style = 'Table Grid'
    h0, h1 = t.rows[0].cells, t.rows[1].cells
    for j, label in ((0, 'Variable'), (1, 'n'), (2, 'Mean difference, ° [min, max]'), (3, 'LoA width, ° [max]')):
        h0[j].merge(h1[j]); h0[j].text = label
    h0[4].merge(h0[6]); h0[4].text = '% of the gait cycle'
    for j, s in zip((4, 5, 6), ('Acceptable', 'Group level only', 'Systematic bias')):
        h1[j].text = s
    for r in rows:
        cells = t.add_row().cells
        vals = [r['variable'], str(r['n']), diff_txt(r), width_txt(r), f"{r['pct_acceptable']:.0f}",
                f"{r['pct_group_only']:.0f}", f"{r['pct_bias']:.0f}"]
        for c, v in zip(cells, vals):
            c.text = v
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.bold = ri < 2 or row.cells[0].text.startswith('Mean of')
    note = d.add_paragraph(footnote)
    for run in note.runs:
        run.font.size = Pt(8)
    out = args.results_dir / f'table_loa_{args.tag}.docx'
    d.save(out)
    print(f'Wrote {out} and the .csv')


if __name__ == '__main__':
    main()
