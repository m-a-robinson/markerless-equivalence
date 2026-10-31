#!/usr/bin/env python3
"""Manuscript table: equivalence outcomes against the SPM result, as a percentage of the gait cycle.

Reads ``results/equivalence_<tag>.npz`` written by ``02_equivalence_analysis.py``.  For every variable each of the 101
gait-cycle nodes has one equivalence outcome and one SPM result, so the table splits the cycle into

                      SPM not significant      SPM significant
    Equivalent                 a                       b
    Inconclusive               c                       d
    Not equivalent             e                       f

(a + b = % equivalent, c + d = % inconclusive, e + f = % not equivalent; a + ... + f = 100 %).  Reading across shows, for
example, how much of the cycle is equivalent although the systems differ significantly (b), inconclusive although SPM
finds no difference (c), or a clear difference beyond the bound (f).  Rows are ordered by % equivalent and the last row is
the mean over the variables.

Outcomes: equivalent = 90 % CI of the mean difference MB - ML inside the bounds; not equivalent = CI entirely beyond a bound;
inconclusive = CI crosses a bound.  SPM = spm1d paired t-test, two-tailed, p < 0.05 (see ``02_equivalence_analysis.py``).

Outputs (``results/``): ``table_equivalence_vs_spm_<tag>.csv`` (all numbers) and, if python-docx is installed,
``table_equivalence_vs_spm_<tag>.docx`` (landscape, with merged headers and a footnote).

Usage
-----
    python 08_summary_table.py [--tag 5deg]
"""
import argparse
import csv
import json
from pathlib import Path

import numpy as np

from common import EQUIVALENT, INCONCLUSIVE, NOT_EQUIVALENT, RESULTS_DIR

OUTCOMES = [('Equivalent', EQUIVALENT), ('Inconclusive', INCONCLUSIVE), ('Not equivalent', NOT_EQUIVALENT)]


def build_rows(z, meta):
    """One dict per variable with the six cross-tab percentages and the totals (percent of the 101 nodes)."""
    rows = []
    for code, title, n, b in zip(meta['variables'], meta['titles'], meta['n'], meta['bounds']):
        cls, sig = z[f'{code}_cls'], z[f'{code}_spm_sig'].astype(bool)
        r = dict(variable=title, n=int(n), bound=float(b), spm_total=100 * sig.mean())
        for name, c in OUTCOMES:
            m = cls == c
            r[f'{name}_total'] = 100 * m.mean()
            r[f'{name}_ns'] = 100 * (m & ~sig).mean()
            r[f'{name}_sig'] = 100 * (m & sig).mean()
        rows.append(r)
    rows.sort(key=lambda r: -r['Equivalent_total'])
    mean = {k: float(np.mean([r[k] for r in rows])) for k in rows[0] if k not in ('variable', 'n', 'bound')}
    rows.append(dict(variable='Mean of the variables', n='', bound=rows[0]['bound'], **mean))
    return rows


def fmt(v):
    return f'{v:.0f}'


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--tag', default='5deg', help='tag of the equivalence results (default 5deg)')
    ap.add_argument('--results-dir', type=Path, default=RESULTS_DIR)
    args = ap.parse_args()

    z = np.load(args.results_dir / f'equivalence_{args.tag}.npz', allow_pickle=False)
    meta = json.loads(str(z['meta']))
    rows = build_rows(z, meta)
    bounds = sorted(set(meta['bounds']))
    bound_txt = f'±{bounds[0]:g}°' if len(bounds) == 1 else 'the variable-specific bounds'

    with open(args.results_dir / f'table_equivalence_vs_spm_{args.tag}.csv', 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(['Variable', 'n'] + [f'{o}_{s}_pct' for o, _ in OUTCOMES for s in ('total', 'SPM_ns', 'SPM_sig')]
                   + ['SPM_sig_total_pct'])
        for r in rows:
            w.writerow([r['variable'], r['n']] + [f"{r[f'{o}_{s}']:.1f}" for o, _ in OUTCOMES for s in ('total', 'ns', 'sig')]
                       + [f"{r['spm_total']:.1f}"])

    footnote = (f'Values are the percentage of the gait cycle (101 points). Equivalent: 90% CI of the mean difference '
                f'(marker-based − markerless) inside {bound_txt}; not equivalent: CI entirely beyond a bound; inconclusive: CI '
                f'crosses a bound. SPM: spm1d paired t-test, two-tailed, p < 0.05; n.s. = not significant, sig. = significant. '
                f'Percentages are rounded, so rows may not sum to 100.')
    print(f'Equivalence outcomes (% of the gait cycle) against the SPM result, bound {bound_txt}')
    head = f"{'Variable':22s}{'n':>3s} | {'Equivalent':>16s} | {'Inconclusive':>16s} | {'Not equivalent':>16s} | SPM sig."
    print(head); print(f"{'':25s} | {'tot  n.s.  sig.':>16s} | {'tot  n.s.  sig.':>16s} | {'tot  n.s.  sig.':>16s} |")
    for r in rows:
        cells = ' | '.join(f"{fmt(r[f'{o}_total']):>4s} {fmt(r[f'{o}_ns']):>4s} {fmt(r[f'{o}_sig']):>4s}   " for o, _ in OUTCOMES)
        print(f"{r['variable']:22s}{str(r['n']):>3s} | {cells}| {fmt(r['spm_total']):>5s}")

    try:
        import docx
        from docx.enum.section import WD_ORIENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
    except ImportError:
        print('python-docx not installed - table written as .csv only')
        return
    d = docx.Document()
    sec = d.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    ncol = 2 + 9 + 1
    t = d.add_table(rows=2, cols=ncol)
    t.style = 'Table Grid'
    h0, h1 = t.rows[0].cells, t.rows[1].cells
    for j, label in ((0, 'Variable'), (1, 'n'), (ncol - 1, 'SPM significant (total)')):
        h0[j].merge(h1[j]); h0[j].text = label
    for g, (name, _) in enumerate(OUTCOMES):
        j = 2 + 3 * g
        h0[j].merge(h0[j + 2]); h0[j].text = name
        for jj, s in zip((j, j + 1, j + 2), ('Total', 'SPM n.s.', 'SPM sig.')):
            h1[jj].text = s
    for r in rows:
        cells = t.add_row().cells
        vals = [r['variable'], str(r['n'])] + [fmt(r[f'{o}_{s}']) for o, _ in OUTCOMES for s in ('total', 'ns', 'sig')] \
               + [fmt(r['spm_total'])]
        for c, v in zip(cells, vals):
            c.text = v
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.bold = ri < 2 or row.cells[0].text.startswith('Mean')
    note = d.add_paragraph(footnote)
    for run in note.runs:
        run.font.size = Pt(8)
    out = args.results_dir / f'table_equivalence_vs_spm_{args.tag}.docx'
    d.save(out)
    print(f'Wrote {out} and the .csv')


if __name__ == '__main__':
    main()
